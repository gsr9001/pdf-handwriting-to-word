"""End-to-end conversion of synthetic PDFs with a mocked AI transcriber."""
import zipfile

import docx
import pytest

import convert_pdf_to_word as conv
from ai_transcriber import TranscriptionError
from conftest import FakeTranscriber


def body_text(path):
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    parts += [c.text for t in d.tables for r in t.rows for c in r.cells]
    return "\n".join(parts)


def assert_no_images(path):
    with zipfile.ZipFile(path) as z:
        assert not [n for n in z.namelist() if n.startswith("word/media/")]
        xml = z.read("word/document.xml").decode()
    assert "<w:drawing" not in xml and "<w:pict" not in xml
    return xml


def test_test_layout_with_ai(test_pdf, tmp_path, quiet):
    fake = FakeTranscriber(responses={
        "question_q02_formula1.png": [("eq", "x=5")],
        "solution_q01_1.png": [("eq", "2+2=4"), ("table", [["a", "b"], ["1", "2"]])],
    })
    r = conv.convert(test_pdf, out_dir=tmp_path / "out", transcriber=fake, log=quiet)
    assert (r["layout"], r["questions"], r["answers"], r["solutions"]) == ("test", 2, 2, 2)
    xml = assert_no_images(r["docx"])
    assert xml.count("<m:oMath>") >= 2
    text = body_text(r["docx"])
    for s in ("Q1.", "Q2.", "What is 2+2?", "Because x = 5", "Solution 1", "Solution 2", "(D)"):
        assert s in text
    d = docx.Document(r["docx"])
    assert [row.cells[1].text.strip() for row in d.tables[0].rows[1:]] == ["4", "5"]  # answer key
    assert len(d.tables) == 2  # answer key + table transcribed from the solution image
    review = r["review"].read_text(encoding="utf-8")
    assert "transcribed by AI (fake-model)" in review
    assert "AI flagged as uncertain — second line smudged" in review
    assert {c[0] for c in fake.calls} == {"question_q02_formula1.png", "solution_q01_1.png"}


def test_manual_transcriptions_take_precedence(test_pdf, tmp_path, quiet):
    module = tmp_path / "manual.py"
    module.write_text(
        'TITLE = "Manual Title"\n'
        'FORMULA_IMAGES = {"question_q02_formula1.png": "x=5"}\n'
        'SOLUTIONS = {1: [("p", "manually typed solution")]}\n'
        'REVIEW_NOTES = {1: "checked by hand"}\n', encoding="utf-8")
    fake = FakeTranscriber()
    r = conv.convert(test_pdf, out_dir=tmp_path / "out", transcriptions=module, transcriber=fake, log=quiet)
    assert fake.calls == []
    text = body_text(r["docx"])
    assert "Manual Title" in text and "manually typed solution" in text
    assert "Solution 1 (PDF page 3): checked by hand" in r["review"].read_text(encoding="utf-8")


def test_without_ai_untranscribed_images_are_flagged(test_pdf, tmp_path, quiet):
    r = conv.convert(test_pdf, out_dir=tmp_path / "out", log=quiet)
    assert_no_images(r["docx"])
    assert "[UNCLEAR IN SOURCE]" in body_text(r["docx"])
    review = r["review"].read_text(encoding="utf-8")
    assert "Solution 1 (PDF page 3): image has NO transcription." in review
    assert r["unclear"] >= 2


def test_ai_failure_is_reported_not_fatal(test_pdf, tmp_path, quiet):
    fake = FakeTranscriber(fail={"solution_q01_1.png"})
    r = conv.convert(test_pdf, out_dir=tmp_path / "out", transcriber=fake, log=quiet)
    review = r["review"].read_text(encoding="utf-8")
    assert "AI transcription failed (simulated API failure)" in review
    assert "[UNCLEAR IN SOURCE]" in body_text(r["docx"])


def test_existing_outputs_are_not_overwritten(test_pdf, tmp_path, quiet):
    out = tmp_path / "out"
    first = conv.convert(test_pdf, out_dir=out, log=quiet)
    stamp = first["docx"].stat().st_mtime_ns
    second = conv.convert(test_pdf, out_dir=out, log=quiet)
    assert second["docx"].name == "SYNTH-TEST_Fully_Editable (2).docx"
    assert second["review"].name == "conversion_review (2).txt"
    assert first["docx"].stat().st_mtime_ns == stamp
    third = conv.convert(test_pdf, out_dir=out, overwrite=True, log=quiet)
    assert third["docx"] == first["docx"]


def test_generic_layout(generic_pdf, tmp_path, quiet):
    fake = FakeTranscriber(default=[("p", "AI page text")])
    r = conv.convert(generic_pdf, out_dir=tmp_path / "out", transcriber=fake, log=quiet)
    assert r["layout"] == "generic"
    assert_no_images(r["docx"])
    text = body_text(r["docx"])
    assert "Lecture notes: soil phases" in text
    assert text.count("AI page text") == 2  # scanned image page + rendered pen-stroke page
    assert [c[1] for c in fake.calls] == ["Page 2 image 1", "Page 3 (rendered page)"]


@pytest.mark.parametrize("make,msg", [
    (lambda p: p / "missing.pdf", "not found"),
    (lambda p: (p / "bad.pdf").write_text("not a pdf") and p / "bad.pdf", "Cannot open"),
])
def test_bad_input_raises(tmp_path, make, msg, quiet):
    with pytest.raises(conv.ConversionError, match=msg):
        conv.convert(make(tmp_path), out_dir=tmp_path / "out", log=quiet)


def test_missing_transcription_module(test_pdf, tmp_path, quiet):
    with pytest.raises(conv.ConversionError, match="Transcription module not found"):
        conv.convert(test_pdf, out_dir=tmp_path, transcriptions=tmp_path / "nope.py", log=quiet)


def test_transcription_error_type_is_shared():
    assert issubclass(TranscriptionError, RuntimeError)


def test_documented_example_module_works(test_pdf, tmp_path, quiet):
    from pathlib import Path
    example = Path(conv.__file__).parent / "examples" / "transcriptions_EXAMPLE.py"
    fake = FakeTranscriber()
    r = conv.convert(test_pdf, out_dir=tmp_path / "out", transcriptions=example, transcriber=fake, log=quiet)
    assert fake.calls == []  # formula and solution both covered by the example module
    text = body_text(r["docx"])
    assert "Test-99 (Synthetic Sample)" in text and "Step A" in text
    assert "value in line 2 is smudged" in r["review"].read_text(encoding="utf-8")


def test_missing_credentials_stop_ai_after_first_image(test_pdf, tmp_path, quiet):
    class NoCredentials(FakeTranscriber):
        def transcribe(self, path, context=""):
            self.calls.append(context)
            raise TranscriptionError("no valid Claude credentials", fatal=True)

    fake = NoCredentials()
    r = conv.convert(test_pdf, out_dir=tmp_path / "out", transcriber=fake, log=quiet)
    assert len(fake.calls) == 1  # not retried for every image
    review = r["review"].read_text(encoding="utf-8")
    assert "AI transcription stopped" in review and "requested but unavailable" in review


def test_keep_images_places_original_pictures(test_pdf, tmp_path, quiet):
    fake = FakeTranscriber()
    r = conv.convert(test_pdf, out_dir=tmp_path / "out", transcriber=fake, log=quiet, keep_images=True)
    assert fake.calls == []  # nothing is sent for transcription
    with zipfile.ZipFile(r["docx"]) as z:
        assert [n for n in z.namelist() if n.startswith("word/media/")]
        assert "<w:drawing" in z.read("word/document.xml").decode()
    assert "Original images are included" in r["review"].read_text(encoding="utf-8")
    assert "[UNCLEAR IN SOURCE] — image" not in body_text(r["docx"])


def test_default_still_embeds_no_images(test_pdf, tmp_path, quiet):
    r = conv.convert(test_pdf, out_dir=tmp_path / "out", transcriber=FakeTranscriber(), log=quiet)
    assert_no_images(r["docx"])

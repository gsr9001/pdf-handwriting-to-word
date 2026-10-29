"""Equation markup -> Office Math, text helpers and table/diagram building."""
from docx import Document

import convert_pdf_to_word as conv


def xml(el):
    from lxml import etree
    return etree.tostring(el, encoding="unicode")


def test_parse_fraction_and_subscripts():
    nodes, _ = conv.parse_math(r"e=\frac{V_{v}}{V_{s}}")
    assert nodes[0] == ("r", "e=", False)
    kind, num, den = nodes[1]
    assert kind == "f"
    assert num == [("sub", [("r", "V", False)], [("r", "v", False)])]
    assert den == [("sub", [("r", "V", False)], [("r", "s", False)])]


def test_subscript_takes_only_last_character_as_base():
    nodes, _ = conv.parse_math("γ_{dry}=2")
    assert nodes[0] == ("sub", [("r", "γ", False)], [("r", "dry", False)])


def test_text_run_is_whole_base_and_subsup():
    nodes, _ = conv.parse_math(r"\text{cm}^{3}")
    assert nodes == [("sup", [("r", "cm", True)], [("r", "3", False)])]
    nodes, _ = conv.parse_math("x_{a}^{2}")
    assert nodes[0][0] == "subsup"


def test_box_and_omml_output():
    el = conv.omath(r"\box{Se=wG}⇒e=\frac{1}{2}")
    out = xml(el)
    for tag in ("m:oMath", "m:borderBox", "m:f", "m:num", "m:den"):
        assert tag in out
    assert "Se=wG" in out


def test_superscript_markers_in_text():
    doc = Document()
    p = doc.add_paragraph()
    conv.add_text(p, f"1 m{conv.SUP_ON}3{conv.SUP_OFF} of soil")
    assert [r.text for r in p.runs] == ["1 m", "3", " of soil"]
    assert p.runs[1].font.superscript is True


def test_markup_text_mixes_inline_equations():
    doc = Document()
    p = doc.add_paragraph()
    conv.add_markup_text(p, "We know $V_{s}=\\frac{V}{1+e}$ always")
    assert "m:oMath" in p._p.xml
    assert p.text.startswith("We know") and p.text.endswith("always")


def test_table_and_diagram_merge():
    doc = Document()
    t = conv.add_table(doc, [["A", "B"], ["full width row", ""], ["1", "2"]])
    assert len(t.rows) == 3
    assert t.rows[1].cells[0]._tc is t.rows[1].cells[1]._tc  # merged
    d = conv.add_diagram(doc, [["Top", "<<"], ["↙", "↘"]])
    assert d.rows[0].cells[0]._tc is d.rows[0].cells[1]._tc


def test_image_block_embeds_a_redrawn_figure(tmp_path):
    from PIL import Image
    fig = tmp_path / "fig.png"
    Image.new("RGB", (400, 200), "white").save(fig)
    doc = Document()
    conv.add_blocks(doc, [("image", str(fig), 3.0)])
    assert len(doc.inline_shapes) == 1
    assert abs(doc.inline_shapes[0].width.inches - 3.0) < 0.01
